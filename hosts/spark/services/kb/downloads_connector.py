#!/usr/bin/env python3
"""downloads_connector.py - KB connector for ~/Documents/Downloads personal docs.

Walks the user's Downloads tree, extracts text from supported document types
(pdf, docx, xlsx, txt) and writes one normalized markdown note per source file
into /var/lib/kb/staging/downloads/, with frontmatter (source path, file type,
ingested timestamp, content hash). The hourly kb-ingest service then embeds the
staging area into pgvector (see default.nix). This script
only touches staging; it never modifies the kb-ingest hot path.

Media (mp4, mov, png, jpg, jpeg, heic, psd, otf) is skipped.

Content-hash dedupe: each note stores the source file's sha256 in frontmatter.
On rerun, if a note for that source already carries the same hash, the file is
left untouched (no re-extract, no rewrite).

PRIVACY DENYLIST (hard, enforced in code - per CLAUDE.md / TOOLS.md):
The directories below are absolutely excluded. We never read, extract, or stage
anything under them. This is enforced by path-prefix exclusion during the walk,
not by convention.
"""

from __future__ import annotations

import hashlib
import os
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path.home() / "Documents" / "Downloads"
STAGING = Path("/var/lib/kb/staging/downloads")

# Hard privacy denylist: absolute directory prefixes that must never be read.
# Enforced in code (see _is_denied). Do not relax without updating CLAUDE.md.
DENYLIST = [
    ROOT / "security",
    ROOT / "documents" / "finance-tax",
    ROOT / "documents" / "travel-identity",
    ROOT / "documents" / "legal-business",
]

# Supported document extensions -> extractor dispatch happens in extract_text.
TEXT_EXTS = {".pdf", ".docx", ".xlsx", ".txt"}

def _is_denied(path: Path) -> bool:
    """True if path is inside any denylisted directory (absolute prefix match)."""
    rp = path.resolve()
    for denied in DENYLIST:
        try:
            rp.relative_to(denied.resolve())
            return True
        except ValueError:
            continue
    return False


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def extract_pdf(path: Path) -> str:
    import fitz  # pymupdf

    parts = []
    with fitz.open(path) as doc:
        for page in doc:
            parts.append(page.get_text())
    return "\n".join(parts)


def extract_docx(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    return "\n".join(p.text for p in d.paragraphs)


def extract_xlsx(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"## {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                parts.append("\t".join(cells))
    wb.close()
    return "\n".join(parts)


def extract_text(path: Path, ext: str) -> str:
    if ext == ".pdf":
        return extract_pdf(path)
    if ext == ".docx":
        return extract_docx(path)
    if ext == ".xlsx":
        return extract_xlsx(path)
    if ext == ".txt":
        return path.read_text(errors="ignore")
    raise ValueError(f"unsupported extension: {ext}")


def note_path(src: Path) -> Path:
    """Deterministic note name keyed on the source path (stable across reruns)."""
    digest = hashlib.sha256(str(src.resolve()).encode()).hexdigest()[:16]
    return STAGING / f"{src.stem}_{digest}.md"


def frontmatter(note: Path) -> dict[str, str]:
    """Read only the generated metadata, never the document body."""
    try:
        with note.open(errors="ignore") as f:
            if f.readline().strip() != "---":
                return {}
            fields = {}
            for _ in range(16):
                line = f.readline(4096)
                if line.strip() == "---":
                    return fields
                key, separator, value = line.partition(":")
                if separator:
                    fields[key] = value.strip()
    except OSError:
        pass
    return {}


def existing_hash(note: Path) -> str | None:
    return frontmatter(note).get("content_hash")


def prune_notes(current: set[Path]) -> int:
    """Remove generated notes whose sources are no longer eligible."""
    removed = 0
    for note in STAGING.glob("*.md"):
        if note in current or note.is_symlink():
            continue
        fields = frontmatter(note)
        source = Path(fields.get("source", ""))
        if (fields.get("source_kind") != "downloads"
                or not source.is_absolute() or not source.is_relative_to(ROOT)
                or not re.fullmatch(r".+_[0-9a-f]{16}\.md", note.name)
                or not re.fullmatch(r"[0-9a-f]{64}", fields.get("content_hash", ""))):
            continue
        note.unlink()
        removed += 1
    return removed


def write_note(src: Path, ext: str, content_hash: str, text: str) -> None:
    STAGING.mkdir(parents=True, exist_ok=True)
    ts = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    note = note_path(src)
    body = (
        "---\n"
        f"source: {src}\n"
        f"file_type: {ext.lstrip('.')}\n"
        f"ingested: {ts}\n"
        f"content_hash: {content_hash}\n"
        "source_kind: downloads\n"
        "---\n\n"
        f"# {src.name}\n\n"
        f"{text.strip()}\n"
    )
    tmp = note.with_suffix(note.suffix + ".tmp")
    tmp.write_text(body)
    tmp.replace(note)


def main() -> int:
    if not ROOT.exists():
        print(f"downloads: {ROOT} does not exist; nothing to do")
        return 0

    written = skipped = denied = unchanged = 0
    current = set()
    walk_errors = []

    for dirpath, dirnames, filenames in os.walk(ROOT, onerror=walk_errors.append):
        d = Path(dirpath)
        # Prune denylisted subtrees in place so we never descend into them.
        dirnames[:] = [
            sub for sub in dirnames if not _is_denied(d / sub)
        ]
        # Skip hidden directories (.git, etc.) for parity with the KB denylist.
        dirnames[:] = [sub for sub in dirnames if not sub.startswith(".")]

        if _is_denied(d):
            continue

        for name in filenames:
            f = d / name
            if name.startswith(".") or f.is_symlink():
                skipped += 1
                continue
            if _is_denied(f):
                denied += 1
                continue
            ext = f.suffix.lower()
            if ext not in TEXT_EXTS:
                skipped += 1
                continue
            current.add(note_path(f))
            try:
                h = sha256_file(f)
            except OSError:
                skipped += 1
                continue
            if existing_hash(note_path(f)) == h:
                unchanged += 1
                continue
            try:
                text = extract_text(f, ext)
            except Exception as e:  # noqa: BLE001 - one bad file shouldn't abort
                print(f"downloads: extract failed for {f}: {e}", file=sys.stderr)
                skipped += 1
                continue
            write_note(f, ext, h, text)
            written += 1

    removed = 0
    if walk_errors:
        print("downloads: incomplete directory walk; retaining existing notes", file=sys.stderr)
    else:
        removed = prune_notes(current)
    print(
        f"downloads: wrote {written}, unchanged {unchanged}, "
        f"skipped {skipped}, denied {denied}, removed {removed} -> {STAGING}"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
